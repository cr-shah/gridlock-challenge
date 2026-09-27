"""Cross-output invariants for the pipeline-published master dataset."""

from __future__ import annotations

import json
import unittest
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

import scoring_engine as scoring


ROOT = Path(__file__).resolve().parents[1]
PUBLISHED = ROOT / "data" / "published"


def read_json(name):
    return json.loads((PUBLISHED / name).read_text(encoding="utf-8"))


class PublishedOutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.master = read_json("gridlock_master_projects.json")
        cls.website = read_json("website_data.json")
        cls.explorer = read_json("project_explorer.json")
        cls.summary = read_json("summary_statistics.json")
        cls.workbook = load_workbook(
            PUBLISHED / "gridlock_all_project_pairs.xlsx", read_only=True, data_only=True
        )

    def test_master_has_exactly_78_unique_projects(self):
        ids = [project["project_id"] for project in self.master["projects"]]
        self.assertEqual(len(ids), 78)
        self.assertEqual(len(set(ids)), 78)
        self.assertEqual(self.master["total_projects"], 78)

    def test_verified_geometry_always_takes_precedence(self):
        for project in self.master["projects"]:
            status = project["geometry_status"]
            if project["verified_geometry"] is not None:
                self.assertEqual(status, "VERIFIED")
                self.assertEqual(project["analysis_geometry"], project["verified_geometry"])
            elif project["estimated_geometry"] is not None:
                self.assertEqual(status, "ESTIMATED")
                self.assertEqual(project["analysis_geometry"], project["estimated_geometry"])
            else:
                self.assertEqual(status, "UNRESOLVED")
                self.assertIsNone(project["analysis_geometry"])

    def test_counts_match_json_and_excel_exports(self):
        statuses = Counter(project["geometry_status"] for project in self.master["projects"])
        self.assertEqual(statuses, Counter(VERIFIED=15, ESTIMATED=44, UNRESOLVED=19))
        self.assertEqual(self.summary["total_projects"], 78)
        self.assertEqual(self.summary["verified_geometry_count"], statuses["VERIFIED"])
        self.assertEqual(self.summary["estimated_geometry_count"], statuses["ESTIMATED"])
        self.assertEqual(self.summary["unresolved_count"], statuses["UNRESOLVED"])

        sheet = self.workbook["Projects"]
        excel_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(excel_rows), 78)
        self.assertEqual(Counter(row[6] for row in excel_rows), statuses)
        self.assertEqual(self.workbook["Within 40 km"].max_row - 1, 38)
        self.assertEqual(self.summary["full_pairs_within_40km"], 38)

    def test_ids_names_and_years_are_identical_across_outputs(self):
        master = {
            project["project_id"]: (
                project["project_name"],
                project["planned_start_year"],
                project["planned_end_year"],
                project["in_service_year"],
            )
            for project in self.master["projects"]
        }
        explorer = {
            project["project_id"]: (
                project["project_name"],
                project["planned_start_year"],
                project["planned_end_year"],
                project["in_service_year"],
            )
            for project in self.explorer["projects"]
        }
        website = {
            project["id"]: (
                project["name"],
                project["planned_start_year"],
                project["planned_end_year"],
                project["in_service_year"],
            )
            for project in self.website["projects"]
        }
        self.assertEqual(master, explorer)
        self.assertEqual(master, website)

        excel = {
            row[0]: (row[2], row[3], row[4], row[5])
            for row in self.workbook["Projects"].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(excel, master)

    def test_website_modes_reference_one_project_catalog(self):
        self.assertEqual(len(self.website["projects"]), 78)
        self.assertNotIn("projects", self.website["modes"]["verified"])
        self.assertNotIn("projects", self.website["modes"]["full"])
        ids = {project["id"] for project in self.website["projects"]}
        for mode in self.website["modes"].values():
            for match in mode["matches"]:
                self.assertIn(match["project_a_id"], ids)
                self.assertIn(match["project_b_id"], ids)
                self.assertNotIn("project_a", match)
                self.assertNotIn("project_b", match)

    def test_scoring_loader_filters_the_master_by_geometry_status(self):
        path = PUBLISHED / "gridlock_master_projects.json"
        verified = scoring.load_projects(path, mode="verified")["projects"]
        full = scoring.load_projects(path, mode="full")["projects"]
        self.assertEqual(len(verified), 15)
        self.assertEqual(len(full), 59)
        self.assertTrue(all(project["geometry_status"] == "VERIFIED" for project in verified))
        self.assertTrue(
            all(project["geometry_status"] in {"VERIFIED", "ESTIMATED"} for project in full)
        )


if __name__ == "__main__":
    unittest.main()
