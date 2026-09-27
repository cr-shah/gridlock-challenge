#!/usr/bin/env python3
"""Generate the analyst workbook from the canonical master; never use Excel as input."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import estimated_coverage as coverage  # noqa: E402
import scoring_engine as scoring  # noqa: E402
from master_dataset import load_master, scoring_projects  # noqa: E402

MASTER_PATH = ROOT / "data/published/gridlock_master_projects.json"
OUTPUT_PATH = ROOT / "data/published/gridlock_all_project_pairs.xlsx"
PAIR_COLUMNS = [
    "DESC Project Name", "DESC Year", "GPC Project Name", "GPC Year", "Year Gap",
    "Distance (km)", "Distance Tier", "DESC Geometry Type", "GPC Geometry Type",
    "DESC Geometry Method", "GPC Geometry Method", "DESC Geometry Confidence",
    "GPC Geometry Confidence",
]
PROJECT_COLUMNS = [
    "Project ID", "Utility", "Project Name", "Planned Start Year", "Planned End Year",
    "In-Service Year", "Geometry Status", "Geometry Type", "Geometry Method",
    "Geometry Confidence",
]
UNRESOLVED_COLUMNS = ["Utility", "Project Name", "Project Year", "Project ID", "Reason"]


def _project_year(project: dict[str, Any]) -> int | str | None:
    start, end = scoring.project_year_window(project)
    if start is None:
        return None
    return int(start) if start == end else f"{int(start)}–{int(end)}"


def _tier(distance: float, intersects: bool) -> str:
    if intersects or distance <= scoring.TOUCHING_EPS_KM:
        return "Touching / Crossing"
    if distance < 1.6:
        return "<1.6 km"
    if distance < 8:
        return "<8 km"
    return "<40 km" if distance < 40 else ">=40 km"


def _pair_rows(projects: list[dict[str, Any]]) -> list[dict[str, Any]]:
    desc = [project for project in projects if project.get("utility") == "DESC"]
    gpc = [project for project in projects if project.get("utility") == "GPC"]
    rows = []
    for a in desc:
        a_parts = coverage.parts_from_geometry(a.get("geometry"), a.get("geometry_type"))
        if not a_parts:
            continue
        for b in gpc:
            b_parts = coverage.parts_from_geometry(b.get("geometry"), b.get("geometry_type"))
            if not b_parts:
                continue
            distance, intersects, _a, _b = coverage.closest_distance_km(a_parts, b_parts)
            _overlap, gap, _years = scoring.timeline_relationship(a, b)
            rows.append({
                "DESC Project Name": a.get("name"),
                "DESC Year": _project_year(a),
                "GPC Project Name": b.get("name"),
                "GPC Year": _project_year(b),
                "Year Gap": gap,
                "Distance (km)": round(distance, 2),
                "Distance Tier": _tier(distance, intersects),
                "DESC Geometry Type": a.get("geometry_type"),
                "GPC Geometry Type": b.get("geometry_type"),
                "DESC Geometry Method": a.get("geometry_method"),
                "GPC Geometry Method": b.get("geometry_method"),
                "DESC Geometry Confidence": a.get("geometry_confidence"),
                "GPC Geometry Confidence": b.get("geometry_confidence"),
                "_distance": distance,
            })
    return sorted(rows, key=lambda row: (
        float("inf") if row["Year Gap"] is None else row["Year Gap"],
        row["_distance"], row["DESC Project Name"], row["GPC Project Name"],
    ))


def _add_sheet(workbook: Workbook, name: str, columns: list[str], rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="123B46")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for index, heading in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 52 if "Name" in heading else 23


def export_workbook(output_path: Path, *, master_path: Path = MASTER_PATH) -> dict[str, int]:
    master = load_master(master_path)
    all_projects = scoring_projects(master, "all")
    mapped = scoring_projects(master, "full")
    pairs = _pair_rows(mapped)
    within = [row for row in pairs if row["_distance"] < 40]
    best = [row for row in within if row["Year Gap"] is not None and row["Year Gap"] <= 2]
    same_year = [row for row in pairs if row["Year Gap"] == 0]
    projects = [{
        "Project ID": project.get("id"), "Utility": project.get("utility"),
        "Project Name": project.get("name"), "Planned Start Year": project.get("planned_start_year"),
        "Planned End Year": project.get("planned_end_year"), "In-Service Year": project.get("in_service_year"),
        "Geometry Status": project.get("geometry_status"), "Geometry Type": project.get("geometry_type"),
        "Geometry Method": project.get("geometry_method"), "Geometry Confidence": project.get("geometry_confidence"),
    } for project in all_projects]
    unresolved = [{
        "Utility": project.get("utility"), "Project Name": project.get("name"),
        "Project Year": _project_year(project), "Project ID": project.get("id"),
        "Reason": project.get("unresolved_reason") or "No verified or estimated geometry available",
    } for project in all_projects if not project.get("geometry")]
    workbook = Workbook()
    workbook.remove(workbook.active)
    _add_sheet(workbook, "Projects", PROJECT_COLUMNS, projects)
    _add_sheet(workbook, "All Evaluated Pairs", PAIR_COLUMNS, pairs)
    _add_sheet(workbook, "Same Year", PAIR_COLUMNS, same_year)
    _add_sheet(workbook, "Within 40 km", PAIR_COLUMNS, within)
    _add_sheet(workbook, "Best Matches", PAIR_COLUMNS, best)
    _add_sheet(workbook, "Unresolved Projects", UNRESOLVED_COLUMNS, unresolved)
    workbook.calculation.fullCalcOnLoad = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {
        "projects": len(projects), "evaluated_pairs": len(pairs),
        "within_40_km": len(within), "best_matches": len(best),
        "unresolved_projects": len(unresolved),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--master", type=Path, default=MASTER_PATH)
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    args = parser.parse_args()
    print(export_workbook(args.output.resolve(), master_path=args.master.resolve()))


if __name__ == "__main__":
    main()
